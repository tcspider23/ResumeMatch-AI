import io
import re
from pathlib import Path
from collections import Counter

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

app = FastAPI(title="ResumeMatch AI API", version="2.0.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

BASE_DIR = Path(__file__).resolve().parent.parent
INDEX_FILE = BASE_DIR / "index.html"

SKILL_GROUPS = {
    "Programming": ["python", "r", "sql", "java", "c++", "c", "javascript", "typescript"],
    "Data & Analytics": ["pandas", "numpy", "scikit-learn", "statistics", "data analysis", "data visualization", "excel", "power bi", "tableau", "time series"],
    "AI & ML": ["machine learning", "deep learning", "nlp", "natural language processing", "tensorflow", "pytorch", "keras", "generative ai", "llm"],
    "Cloud & Big Data": ["aws", "azure", "gcp", "hadoop", "spark", "docker"],
    "Web & Backend": ["html", "css", "flask", "fastapi", "django", "mongodb", "mysql", "postgresql"],
    "Professional": ["communication", "leadership", "problem solving", "git", "github"],
}
SKILLS = [s for group in SKILL_GROUPS.values() for s in group]
STOP = set("a an the and or but for with from into of to in on at by is are was were be been this that as it we you your our their they he she i me my".split())


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text.lower()).strip()


def extract_skills(text: str):
    low = clean_text(text)
    return sorted({skill for skill in SKILLS if re.search(r"(?<!\w)" + re.escape(skill) + r"(?!\w)", low)})


def grouped_skills(text: str):
    found = set(extract_skills(text))
    return {group: sorted(found.intersection(items)) for group, items in SKILL_GROUPS.items() if found.intersection(items)}


def top_keywords(text: str, limit=12):
    words = re.findall(r"[a-zA-Z][a-zA-Z+#.-]{2,}", clean_text(text))
    counts = Counter(w for w in words if w not in STOP and not w.isdigit())
    return [w for w, _ in counts.most_common(limit)]


def extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
        if text.strip():
            return text
    except Exception:
        pass
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception:
        return ""


def extract_docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_resume(name: str, data: bytes) -> str:
    name = (name or "").lower()
    if name.endswith(".pdf"): return extract_pdf(data)
    if name.endswith(".docx"): return extract_docx(data)
    if name.endswith(".txt"): return data.decode("utf-8", errors="ignore")
    return ""


def make_summary(resume: str, matched: list[str]):
    title = "Data and technology professional"
    low = clean_text(resume)
    for candidate in ["data scientist", "data analyst", "machine learning engineer", "software engineer", "developer"]:
        if candidate in low:
            title = candidate.title()
            break
    skills = ", ".join(x.title() for x in matched[:5]) or "relevant technical skills"
    return f"{title} with experience and interest in {skills}. Skilled at applying analytical and technology-driven approaches to solve practical problems and deliver clear, useful results."


def analyze(resume: str, job: str):
    resume_clean, job_clean = clean_text(resume), clean_text(job)
    try:
        matrix = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True).fit_transform([resume_clean, job_clean])
        similarity = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0] * 100)
    except ValueError:
        similarity = 0.0

    resume_skills, job_skills = set(extract_skills(resume)), set(extract_skills(job))
    matched = sorted(resume_skills & job_skills)
    missing = sorted(job_skills - resume_skills)
    skill_coverage = (len(matched) / len(job_skills) * 100) if job_skills else 0.0
    keyword_set = set(top_keywords(job, 18))
    resume_words = set(re.findall(r"[a-zA-Z][a-zA-Z+#.-]{2,}", resume_clean))
    keyword_coverage = (len(keyword_set & resume_words) / len(keyword_set) * 100) if keyword_set else 0.0
    ats = min(100.0, 0.55 * skill_coverage + 0.30 * keyword_coverage + 0.15 * similarity)
    readiness = min(100.0, 0.60 * similarity + 0.40 * skill_coverage)

    suggestions = []
    for skill in missing[:6]:
        suggestions.append(f"If you genuinely have {skill.title()} experience, make it visible in your Skills or Projects section.")
    if keyword_coverage < 60:
        suggestions.append("Mirror important job-description terminology naturally in relevant experience or project bullets.")
    if len(resume.split()) < 180:
        suggestions.append("Add measurable project or experience details so recruiters can see your impact.")
    if not suggestions:
        suggestions.append("Your resume covers the detected requirements well. Keep achievements specific and measurable.")

    return {
        "score": round(similarity, 1),
        "ats_score": round(ats, 1),
        "readiness": round(readiness, 1),
        "skill_coverage": round(skill_coverage, 1),
        "keyword_coverage": round(keyword_coverage, 1),
        "matched": matched,
        "missing": missing,
        "resume_skills": sorted(resume_skills),
        "job_skills": sorted(job_skills),
        "skill_groups": grouped_skills(resume),
        "job_skill_groups": grouped_skills(job),
        "job_keywords": top_keywords(job, 12),
        "suggestions": suggestions,
        "summary": make_summary(resume, matched),
        "method": ["Text extraction", "Preprocessing", "TF-IDF vectorization", "Cosine similarity", "Skill extraction", "ATS-style keyword analysis"],
    }


@app.get("/")
def home():
    return FileResponse(INDEX_FILE)


@app.get("/api/health")
def health():
    return {"status": "ok", "service": "ResumeMatch AI", "version": "2.0.0"}


@app.post("/api/analyze")
async def analyze_resume(resume: UploadFile = File(...), job_description: str = Form(...)):
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Please provide a job description.")
    text = extract_resume(resume.filename or "", await resume.read())
    if not text.strip():
        raise HTTPException(status_code=400, detail="Could not extract text. Use a text-based PDF, DOCX or TXT file.")
    result = analyze(text, job_description)
    result["resume_text"] = text
    return result


@app.post("/api/chat")
async def chat(question: str = Form(...), resume_text: str = Form(""), job_description: str = Form("")):
    q = clean_text(question)
    result = analyze(resume_text, job_description) if resume_text and job_description else None
    if not q:
        return {"answer": "Ask me about your resume, skills, match score, ATS readiness, or improvement areas."}
    if result:
        if "missing" in q or "skill gap" in q:
            answer = "Your main detected skill gaps are: " + (", ".join(x.title() for x in result["missing"][:8]) or "none detected") + "."
        elif "score" in q or "match" in q:
            answer = f"Your current NLP match score is {result['score']:.1f}%, with ATS-style readiness at {result['ats_score']:.1f}%."
        elif "improve" in q or "better" in q:
            answer = "Start with these improvements: " + " ".join(result["suggestions"][:3])
        elif "summary" in q:
            answer = result["summary"]
        elif "keyword" in q:
            answer = "Important job keywords detected: " + ", ".join(x.title() for x in result["job_keywords"]) + "."
        else:
            answer = f"I found {len(result['matched'])} matched skills and {len(result['missing'])} missing skills. Ask me about your score, skill gap, keywords, summary, or improvements."
    else:
        answer = "Upload a resume and add a job description in Analyze Resume first, then I can answer with project-specific insights."
    return {"answer": answer}
